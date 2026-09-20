from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION

OUT = r"docs/붙임4_제안요약서_목적필요성보완.docx"

NAVY = "1F4E79"
BLUE = "D9EAF7"
PALE = "F4F7FA"
GRAY = "6B7280"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margin(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_font(run, size=10, bold=False, color=None):
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def format_paragraph(p, before=0, after=5, line=1.45, align=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align


def add_text(doc, text, bold=False, size=10.2, color=None, before=0, after=5, line=1.45):
    p = doc.add_paragraph()
    format_paragraph(p, before, after, line)
    r = p.add_run(text)
    set_font(r, size, bold, color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    format_paragraph(p, before=10 if level == 1 else 6, after=6, line=1.1)
    r = p.add_run(text)
    set_font(r, 14 if level == 1 else 11.5, True, NAVY)
    p.paragraph_format.keep_with_next = True
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    format_paragraph(p, after=3, line=1.35)
    r = p.add_run(text)
    set_font(r, 9.8)
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, NAVY)
        set_cell_margin(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        format_paragraph(p, after=0, line=1.1, align=WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(header)
        set_font(r, 9.2, True, "FFFFFF")
        if widths:
            cell.width = Cm(widths[i])
    for row_i, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            set_cell_margin(cell)
            if row_i % 2 == 1:
                set_cell_shading(cell, PALE)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            format_paragraph(p, after=0, line=1.25)
            r = p.add_run(value)
            set_font(r, 8.8)
            if widths:
                cell.width = Cm(widths[i])
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_page_number(footer):
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("- ")
    set_font(r, 8, color=GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    p._p.append(fld)
    r = p.add_run(" -")
    set_font(r, 8, color=GRAY)


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.7)
section.bottom_margin = Cm(1.5)
section.left_margin = Cm(1.8)
section.right_margin = Cm(1.8)

styles = doc.styles
styles["Normal"].font.name = "Malgun Gothic"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
styles["Normal"].font.size = Pt(10.2)
styles["Title"].font.name = "Malgun Gothic"
styles["Title"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")

add_page_number(section.footer)

# Page 1
p = doc.add_paragraph()
format_paragraph(p, after=7, align=WD_ALIGN_PARAGRAPH.CENTER)
r = p.add_run("붙임 4  제안 요약서")
set_font(r, 11, True, GRAY)

p = doc.add_paragraph(style="Title")
format_paragraph(p, before=18, after=10, line=1.2, align=WD_ALIGN_PARAGRAPH.CENTER)
r = p.add_run("대구 음식점 소상공인을 위한\n공공데이터 기반 AI 주간 운영 가이드")
set_font(r, 22, True, NAVY)

p = doc.add_paragraph()
format_paragraph(p, after=22, align=WD_ALIGN_PARAGRAPH.CENTER)
r = p.add_run("장사메이트")
set_font(r, 13, True, GRAY)

info = doc.add_table(rows=2, cols=4)
info.style = "Table Grid"
info.alignment = WD_TABLE_ALIGNMENT.CENTER
for row in info.rows:
    for c in row.cells:
        set_cell_margin(c)
for i, label in enumerate(["접수번호", "", "팀명", "노드"]):
    p = info.rows[0].cells[i].paragraphs[0]
    format_paragraph(p, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(label)
    set_font(r, 9.5, i in (0, 2), "FFFFFF" if i in (0, 2) else None)
    if i in (0, 2): set_cell_shading(info.rows[0].cells[i], NAVY)
for i, label in enumerate(["", "", "제안 분야", "AI 기반 공공데이터 활용"]):
    p = info.rows[1].cells[i].paragraphs[0]
    format_paragraph(p, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(label)
    set_font(r, 9.2, i == 2, "FFFFFF" if i == 2 else None)
    if i == 2: set_cell_shading(info.rows[1].cells[i], NAVY)

add_heading(doc, "과제 요약", 1)
add_text(doc, "대구 음식점 사장님은 매주 달라지는 날씨, 공휴일, 지역 행사, 식자재 가격과 주변 영업환경을 여러 경로에서 확인한 뒤 발주·인력·배달 준비를 결정해야 한다. 장사메이트는 상호명과 주소, 대표 메뉴를 입력하면 가게의 위치와 메뉴 특성에 맞춰 이번 주에 확인할 운영 항목을 정리해 주는 실사용형 웹서비스다.")
add_text(doc, "식품 인허가, 기상청 예보, KAMIS 농산물 가격, TourAPI 지역 행사, 천문연구원 공휴일 정보를 결합하고, 반경 500m 안의 음식점 수와 유사업종 수를 함께 제공한다. 서비스는 매출·이익·현금흐름을 예측하지 않으며, 검증 가능한 공공데이터를 근거로 최대 3개의 운영 참고사항과 근거보기를 제공한다. 챗봇은 해당 리포트 안의 데이터만 설명하며, 산출 근거가 없는 구체적 매출·손익 수치 질문에는 답하지 않는다.")

add_heading(doc, "1  제안내용의 목적 및 필요성", 1)
add_heading(doc, "문제 정의", 2)
add_text(doc, "외식업 사장님에게 날씨와 행사 일정, 식자재 가격의 변화는 발주와 인력 배치, 포장·배달 준비에 직접 연결되는 정보다. 하지만 이 정보는 기관별로 흩어져 있고, 가게의 위치·업종과 연결해 이번 주에 무엇을 준비해야 하는지 판단하기 어렵다. 통계청 2022년 소상공인실태조사에서 숙박·음식점업 경영 애로사항은 원재료비 60.2%, 상권 쇠퇴 39.3%, 경쟁 심화 32.7% 순으로 나타났다.")
add_heading(doc, "핵심 솔루션", 2)
add_text(doc, "장사메이트는 실제 가게의 위치와 대표 메뉴를 기준으로 날씨·행사·공휴일·식자재 가격·주변 음식점 정보를 한 리포트에 결합한다. 사장님은 여러 공공 사이트를 오가며 정보를 해석하는 대신, 이번 주에 확인할 운영 항목과 근거를 한 흐름에서 볼 수 있다.")
add_heading(doc, "주요 기능", 2)
add_text(doc, "① 대구 음식점 인허가 데이터 기반 가게 검색과 메뉴 카테고리 분류  ② 가게별 이번 주 운영 가이드 최대 3개와 근거보기  ③ 최근 7일 식자재 가격 그래프, 반경 500m 내 음식점·유사업종 수, 반경 3km 내 행사 정보  ④ 리포트 안의 사실만 설명하는 AI 챗봇과 오늘 기준 재분석 기능", size=9.7, after=4)
add_text(doc, "출처: 통계청 「2022년 소상공인실태조사」", size=8.5, color=GRAY, after=0)

doc.add_page_break()

# Page 2
add_heading(doc, "2  서비스 구조 및 기술", 1)
add_text(doc, "서비스는 실존 음식점 인허가 정보로 가게를 찾고, 가게 위치와 메뉴 카테고리를 기준으로 공공데이터를 수집·정리한다. 리포트에는 데이터 출처와 수집 기준일을 함께 제시해, 사용자가 운영 판단에 사용된 정보의 범위를 확인할 수 있게 한다.")

add_heading(doc, "서비스 흐름", 2)
add_table(doc, ["1. 가게 등록", "2. 데이터 결합", "3. 주간 운영 가이드", "4. 재분석·질문"], [[
    "상호명·주소 검색\n대표 메뉴 입력",
    "날씨·행사·가격·공휴일\n주변 음식점 현황",
    "이번 주 확인할 항목\n최대 3개와 근거보기",
    "오늘 기준 재분석\n리포트 근거 챗봇",
]], [4.0, 4.0, 4.0, 4.0])

add_heading(doc, "공공데이터와 서비스 적용", 2)
add_table(doc, ["데이터", "서비스에서의 활용", "표시 기준"], [
    ["대구 식품 인허가", "가게 검색, 업태·주소·좌표 확인, 주변 음식점 집계", "검색 대상 24,079곳"],
    ["기상청 단기예보", "이번 주 날씨 변화와 운영 점검 항목 연결", "예보 발표 시점 표시"],
    ["KAMIS 농산물 가격", "메뉴 특성에 맞는 식자재 가격과 최근 7일 추이", "품목·조사 기준일 표시"],
    ["TourAPI 지역 행사", "가게 반경 3km 내 예정 행사 확인", "행사명·일정·거리 기준 표시"],
    ["공휴일·주변 업소", "공휴일 일정, 500m 내 음식점·유사업종 수", "공공데이터 기준"],
], [3.2, 8.7, 4.1])

add_heading(doc, "기술 스택", 2)
add_table(doc, ["구분", "적용 기술", "역할"], [
    ["Frontend", "Next.js 16, TypeScript, Axios, React Hook Form, Zod", "가게 등록, 리포트, 근거보기, 챗봇 화면"],
    ["Backend", "Java 21, Spring Boot 3.5, Spring Data JPA, Spring Validation", "공공데이터 Provider, 리포트 조합, API 제공"],
    ["데이터베이스", "H2 개발 DB, MySQL profile 준비", "개발·운영 환경 분리 준비"],
    ["AI", "Groq 기반 챗봇, 리포트 사실 범위 검증", "수치·날짜 임의 생성 방지"],
], [2.5, 7.4, 6.1])

doc.add_page_break()

# Page 3
add_heading(doc, "3  주요 기능", 1)
add_heading(doc, "가게 등록과 메뉴 기반 개인화", 2)
add_text(doc, "사장님은 대구 음식점 인허가 데이터에서 자신의 가게를 검색해 선택한다. 직접 주소를 입력할 수도 있으며, 대표 메뉴는 입력 내용에 따라 카테고리로 분류한다. 업종이 명확한 메뉴 후보를 우선 제안해 업종과 맞지 않는 메뉴가 추천되는 문제를 줄인다.")

add_heading(doc, "이번 주 운영 가이드와 근거보기", 2)
add_text(doc, "리포트는 현재 가게의 메뉴 카테고리, 위치, 날짜를 기준으로 이번 주에 확인할 운영 항목을 최대 3개로 압축한다. 권고 문구에는 근거 수치나 긴 출처를 반복해 노출하지 않고, ‘근거보기’에서 날씨·행사·가격·주변 현황의 원자료와 데이터 상태를 확인할 수 있다.")
add_table(doc, ["화면", "사용자에게 제공하는 내용"], [
    ["이번 주 처방", "날씨, 행사, 식자재 가격, 공휴일 등을 바탕으로 발주·인력·배달·포장 준비 점검"],
    ["식자재 가격", "대표 메뉴와 연결된 품목의 최근 7일 가격을 꺾은선 그래프로 표시"],
    ["주변 현황", "반경 500m 내 음식점 수와 유사업종 수를 제공. 경쟁강도 등급은 사용하지 않음"],
    ["행사 정보", "반경 3km 내 예정 행사명과 날짜만 표시. 정보가 없으면 없음을 명확히 안내"],
    ["처방 이력", "브라우저에 저장된 이전 리포트와 현재 리포트를 구분해 확인"],
], [3.4, 12.6])

add_heading(doc, "리포트 기반 AI 챗봇", 2)
add_text(doc, "챗봇은 선택한 리포트에 포함된 사실만 바탕으로 답한다. 날씨·행사·가격·주변 현황의 의미나 준비 방법은 설명할 수 있지만, 현재 서비스가 산출하지 않는 매출·이익·고객 수·발주량·효과의 구체적 수치는 답하지 않도록 제한한다. 외부 AI 응답 오류나 API 미연결 시에는 그 상태를 화면에 명확히 표시한다.")

add_heading(doc, "재분석과 데이터 최신성", 2)
add_text(doc, "사용자는 ‘오늘 기준으로 다시 분석’을 눌러 새 리포트를 생성할 수 있다. 최신 예보와 수집된 가격·행사·상권 정보를 다시 반영하고, 각 정보의 출처와 기준일을 리포트에서 확인할 수 있다. 데이터가 확인되지 않는 경우에는 근거 없는 권고를 만들지 않고 해당 정보를 제외하거나 이용 불가 상태를 안내한다.")

doc.add_page_break()

# Page 4
add_heading(doc, "4  제안내용의 차별성", 1)
add_heading(doc, "가게 단위의 정보 결합", 2)
add_text(doc, "기존 상권 정보는 지역 통계나 단일 지표를 확인하는 데 그치는 경우가 많다. 장사메이트는 실제 가게의 주소와 메뉴 카테고리를 기준으로 날씨, 행사, 식자재 가격, 공휴일, 주변 음식점 정보를 한 리포트에 모은다. 사용자는 여러 사이트를 오가며 정보를 해석하는 대신, 가게에 필요한 이번 주 확인 항목부터 볼 수 있다.")

add_heading(doc, "예측이 아닌 검증 가능한 운영 참고", 2)
add_text(doc, "서비스는 매출액, 매출 증감률, 이익, 현금흐름, 필요 발주량을 예측하지 않는다. 공공데이터에서 확인된 사실과 데이터 상태를 바탕으로 ‘준비 권장’, ‘점검 권장’ 수준의 운영 참고사항을 제공한다. 이 범위를 명확히 해 근거 없는 금융·성과 예측으로 오해될 여지를 줄인다.")

add_heading(doc, "조언과 근거를 분리한 신뢰 구조", 2)
add_text(doc, "첫 화면에서는 사장님이 바로 읽을 수 있는 짧은 운영 가이드를 제시하고, 필요한 경우 근거보기에서 출처·기준일·원자료를 확인하도록 구성했다. 권고마다 sourceIds를 연결하고, 검증에 실패한 조언은 표시하지 않는다. 챗봇도 리포트에 없는 숫자나 날짜를 새로 만들지 않도록 별도의 검증 단계를 둔다.")

add_heading(doc, "현재 기능과 확장 계획의 구분", 2)
add_table(doc, ["현재 구현", "향후 검토 가능한 확장"], [[
    "공공데이터 기반 주간 운영 가이드\n가게·메뉴별 데이터 결합\n근거보기와 리포트 범위 챗봇\n재분석과 데이터 상태 표시",
    "사장님 동의 기반 POS·주문·금융 데이터 연계\n충분한 검증 후 비용·현금흐름 분석 보조\n정책자금·지역 지원사업 정보 연계\n지역·업종 확대",
]], [8.0, 8.0])
add_text(doc, "금융거래 기반 현금흐름 예측이나 자금 부족 조기경보는 현재 서비스 기능이 아니다. 향후 데이터 동의, 개인정보 보호, 충분한 검증 체계가 갖춰진 뒤 별도 과제로 검토한다.", size=9.2, color=GRAY, after=0)

doc.add_page_break()

# Page 5
add_heading(doc, "5  기대효과 사업화 및 향후 계획", 1)
add_heading(doc, "소상공인", 2)
add_bullet(doc, "날씨·행사·식자재 가격·공휴일·주변 현황을 한 화면에서 확인해 정보 탐색 시간을 줄일 수 있다.")
add_bullet(doc, "매출을 단정적으로 예측하지 않고, 이번 주 발주·인력·배달·포장 준비에서 확인할 항목을 빠르게 점검할 수 있다.")
add_bullet(doc, "근거보기와 데이터 상태 표시를 통해 어떤 데이터가 현재 리포트에 사용됐는지 직접 확인할 수 있다.")

add_heading(doc, "지역사회와 공공데이터", 2)
add_bullet(doc, "흩어진 공공데이터를 음식점 사장님의 일상적인 운영 판단에 연결해 데이터 활용 접근성을 높인다.")
add_bullet(doc, "지역 행사와 생활권 음식점 정보를 가게 위치 기준으로 제시해 지역 정보의 현장 활용 가능성을 넓힌다.")
add_bullet(doc, "향후 이용 패턴과 데이터 품질을 검증하면 지역별·업종별 지원정책 연구를 위한 보조 지표로 발전시킬 여지가 있다.")

add_heading(doc, "평가항목별 제안의 강점", 2)
add_table(doc, ["평가 항목", "제안의 대응"], [
    ["실용성 및 문제해결력", "가게 등록부터 이번 주 점검 항목 확인까지 한 흐름으로 제공한다. 여러 공공 사이트를 오가는 대신, 실제 음식점 위치·메뉴 기준으로 발주·인력·배달 준비에 필요한 정보만 정리한다."],
    ["기술 완성도", "Next.js와 Spring Boot 기반의 분리된 구조, 데이터 Provider 인터페이스, 입력 검증, 근거 연결, 리포트 범위 챗봇을 구현했다. 데이터 오류 시에는 근거 없는 권고를 표시하지 않는다."],
    ["사업화 확장 가능성", "음식점 사장님 대상 기본 운영 가이드에서 출발해, 지역 소상공인 지원기관·상권 활성화 사업과의 B2B 협력, 업종·지역 확장, 동의 기반 데이터 연계를 단계적으로 검토할 수 있다."],
    ["제안 충실도", "문제 정의, 실제 데이터 흐름, 기능 범위, 검증 방식, 기대효과와 향후 단계의 경계를 명확히 제시한다. 현재 구현 기능과 미래 확장 기능을 구분해 과장 가능성을 줄였다."],
], [3.5, 12.5])

add_heading(doc, "사업화와 단계별 고도화 계획", 2)
add_table(doc, ["단계", "내용", "검증 기준"], [
    ["1단계 MVP", "대구 음식점 대상 가게 등록, 공공데이터 결합, 주간 운영 가이드 제공", "데이터 상태 표시, 근거보기, 전체 사용자 흐름 점검"],
    ["2단계 사업화", "지역 소상공인 지원기관·상권 활성화 사업과 연계한 도입 검토, 업종별 운영 가이드 고도화", "사용자 피드백, 재방문·재분석 이용 지표, 데이터 누락률"],
    ["3단계 연계 검토", "동의 기반의 주문·POS·금융 데이터 연계 가능성 검토", "개인정보 보호, 데이터 동의, 별도 성능 검증과 책임 기준"],
], [2.7, 7.0, 6.3])

add_heading(doc, "제안의 범위", 2)
add_text(doc, "장사메이트는 ‘이번 주 가게 운영에 필요한 변화를 공공데이터로 확인하고 준비하게 돕는 서비스’다. 현재 MVP는 데이터 기반 운영 참고와 설명 가능성에 집중한다. 매출·이익·현금흐름 예측이나 금융상품 추천은 서비스 범위에 포함하지 않으며, 향후 별도 검증과 사용자 동의를 전제로 검토한다.")

add_heading(doc, "주요 데이터 출처", 2)
add_text(doc, "식품의약품안전처 식품위생업소 인허가, 기상청 단기예보, 한국농수산식품유통공사 KAMIS, 한국관광공사 TourAPI, 한국천문연구원 특일 정보, 통계청 「2022년 소상공인실태조사」", size=9.3, after=0)

doc.core_properties.title = "대구 음식점 소상공인을 위한 공공데이터 기반 AI 주간 운영 가이드"
doc.core_properties.subject = "붙임 4 제안 요약서 수정본"
doc.core_properties.author = "노드"
doc.save(OUT)
print(OUT)
